from __future__ import annotations

import json
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from zoneinfo import ZoneInfo

from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, urlparse


ROOT = Path(__file__).resolve().parent
SOURCE_XLSX = Path("/Users/yu-kyoungin/Downloads/01.2026년 CATV 계획점검표.xlsx")
OUTPUT_DIR = ROOT / "outputs"
WORKBOOK_PATH = OUTPUT_DIR / "CATV_계획점검_웹입력_작업본.xlsx"

PLAN_SHEET = "계획점검 이행일"
FIBER_SHEET = "CA_광선로 측정_2026"
MEASURE_SHEETS = {"K1": "K1 측정값", "K2": "K2 측정값", "사외": "사외 측정값"}
KST = ZoneInfo("Asia/Seoul")
NAME_ALIASES = {
    "나노파크": ["나노파크 & 스포렉스", "나노파크&스포렉스"],
    "세미콘프라자": ["세미콘 프라자"],
}
FORCED_MEASURE_BLOCKS = {
    "해피홈기숙사": {"sheet": "사외 측정값", "headerRow": 37, "title": "해피홈기숙사"},
}

SIGNAL_STEPS = [
    [3],
    [4],
    [7],
    [10],
    [13],
    [19],
    [25],
    [26, 27],
    [20, 21],
    [14, 15],
    [11, 12],
    [8, 9],
    [5, 6],
    [28, 30],
]


def ensure_workbook() -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)
    if not WORKBOOK_PATH.exists():
        shutil.copy2(SOURCE_XLSX, WORKBOOK_PATH)


def load_workbook():
    ensure_workbook()
    return openpyxl.load_workbook(WORKBOOK_PATH)


def today() -> datetime:
    return datetime.now(KST)


def norm(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").strip()


def match_key(value: Any) -> str:
    text = norm(value)
    return re.sub(r"[\s()&・ㆍ/._-]+", "", text).lower()


def name_candidates(building: str) -> list[str]:
    return [building, *NAME_ALIASES.get(building, [])]


def names_match(left: Any, right: Any) -> bool:
    left_key = match_key(left)
    right_key = match_key(right)
    if not left_key or not right_key:
        return False
    return left_key == right_key


def active_region(ws, row: int) -> str:
    value = None
    for r in range(4, row + 1):
        current = ws.cell(r, 2).value
        if current:
            value = current
    return norm(value)


def find_plan_row(ws, building: str) -> int:
    for row in range(4, 30):
        if norm(ws.cell(row, 3).value) == building:
            return row
    raise ValueError(f"계획점검 이행일 C4:C29에서 '{building}'을 찾을 수 없습니다.")


def find_month_column(ws, month: int) -> int:
    target = f"{month}월"
    for col in range(1, ws.max_column + 1):
        if norm(ws.cell(3, col).value) == target:
            return col
    raise ValueError(f"계획점검 이행일 3행에서 '{target}'을 찾을 수 없습니다.")


def find_month_row(ws, month: int) -> int:
    target = f"{month}월"
    for row in range(1, ws.max_row + 1):
        if norm(ws.cell(row, 2).value) == target:
            return row
    raise ValueError(f"{ws.title} B열에서 '{target}'을 찾을 수 없습니다.")


def find_month_row_in_block(ws, header_row: int, month: int) -> int:
    target = f"{month}월"
    for row in range(header_row + 3, min(header_row + 15, ws.max_row + 1)):
        if norm(ws.cell(row, 2).value) == target:
            return row
    raise ValueError(f"{ws.title} {header_row}행 블록에서 '{target}'을 찾을 수 없습니다.")


def find_building_row(ws, building: str, col: int = 3) -> int:
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, col).value
        if any(names_match(value, candidate) for candidate in name_candidates(building)):
            return row
    raise ValueError(f"{ws.title} {get_column_letter(col)}열에서 '{building}'을 찾을 수 없습니다.")


def find_sheet_name(wb, building: str) -> str | None:
    for candidate in name_candidates(building):
        if candidate in wb.sheetnames:
            return candidate
    for sheet_name in wb.sheetnames:
        if names_match(sheet_name, building):
            return sheet_name
    return None


def measure_blocks(ws) -> list[dict[str, Any]]:
    blocks = []
    for row in range(1, ws.max_row + 1):
        if norm(ws.cell(row, 2).value).startswith("채널"):
            title = display_value(ws.cell(row - 1, 2).value) if row > 1 else None
            blocks.append({"headerRow": row, "titleRow": row - 1, "title": norm(title)})
    return blocks


def find_measure_block(wb, region: str, building: str) -> dict[str, Any] | None:
    forced = FORCED_MEASURE_BLOCKS.get(building)
    if forced and forced["sheet"] in wb.sheetnames:
        return {**forced, "region": region}

    sheet_name = MEASURE_SHEETS.get(region)
    if not sheet_name or sheet_name not in wb.sheetnames:
        return None
    ws = wb[sheet_name]
    for block in measure_blocks(ws):
        if any(names_match(block["title"], candidate) for candidate in name_candidates(building)):
            return {**block, "sheet": sheet_name, "region": region}
    return None


def display_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return value


def numeric_or_text(value: Any) -> Any:
    if value is None or value == "":
        return "-"
    text = str(value).strip()
    try:
        return float(text)
    except ValueError:
        return text


def sectors_payload(wb) -> list[dict[str, Any]]:
    ws = wb[PLAN_SHEET]
    rows = []
    for row in range(4, 30):
        building = norm(ws.cell(row, 3).value)
        if building:
            rows.append({"row": row, "building": building, "region": active_region(ws, row)})
    return rows


def plan_history_payload(wb) -> dict[str, Any]:
    ws = wb[PLAN_SHEET]
    current_month = today().month
    months = [{"month": idx, "label": norm(ws.cell(3, idx + 3).value)} for idx in range(1, 13)]
    rows = []
    month_counts = {str(idx): 0 for idx in range(1, 13)}
    current_month_checked = 0
    for row in range(4, 30):
        building = norm(ws.cell(row, 3).value)
        if not building:
            continue
        checks = {}
        checked_months = 0
        for idx in range(1, 13):
            value = display_value(ws.cell(row, idx + 3).value)
            checks[str(idx)] = value
            if value not in (None, ""):
                checked_months += 1
                month_counts[str(idx)] += 1
        if checks[str(current_month)] not in (None, ""):
            current_month_checked += 1
        rows.append(
            {
                "row": row,
                "region": active_region(ws, row),
                "building": building,
                "checks": checks,
                "checkedMonths": checked_months,
            }
        )
    total = len(rows)
    return {
        "months": months,
        "rows": rows,
        "summary": {
            "totalBuildings": total,
            "currentMonth": current_month,
            "checkedBuildings": current_month_checked,
            "uncheckedBuildings": total - current_month_checked,
            "monthCounts": month_counts,
        },
    }


def fiber_headers(ws) -> list[dict[str, Any]]:
    headers = []
    for col in range(8, 11):
        headers.append(
            {
                "col": col,
                "letter": get_column_letter(col),
                "label": norm(ws.cell(3, col).value),
            }
        )
    return headers


def signal_headers(ws, cols: list[int], header_row: int = 3) -> list[dict[str, Any]]:
    headers = []
    for col in cols:
        headers.append(
            {
                "col": col,
                "letter": get_column_letter(col),
                "channel": norm(ws.cell(header_row, col).value),
                "measure": norm(ws.cell(header_row + 1, col).value),
                "kind": norm(ws.cell(header_row + 2, col).value),
            }
        )
    return headers


def previous_nonblank(ws, before_row: int, col: int) -> Any:
    for row in range(before_row - 1, 5, -1):
        value = ws.cell(row, col).value
        if value not in (None, ""):
            return value
    return None


def previous_nonblank_in_block(ws, header_row: int, before_row: int, col: int) -> Any:
    for row in range(before_row - 1, header_row + 2, -1):
        value = ws.cell(row, col).value
        if value not in (None, ""):
            return value
    return None


def grouped_measure_values(ws, row: int, header_row: int) -> list[dict[str, Any]]:
    groups = []
    current = None
    for col in range(3, 28):
        channel = norm(ws.cell(header_row, col).value)
        if channel:
            current = {
                "channel": channel,
                "rfPower": None,
                "evm": None,
                "mer": None,
                "values": [],
            }
            groups.append(current)
        if not current:
            continue
        value = ws.cell(row, col).value
        measure = norm(ws.cell(header_row + 1, col).value)
        kind = norm(ws.cell(header_row + 2, col).value)
        if value in (None, ""):
            continue
        label = kind or measure
        item = {
            "label": label,
            "measure": measure,
            "kind": kind,
            "value": display_value(value),
        }
        current["values"].append(item)
        if kind == "신호값" or "RF Power" in measure:
            current["rfPower"] = display_value(value)
        elif kind == "EVM":
            current["evm"] = display_value(value)
        elif kind == "MER":
            current["mer"] = display_value(value)
    return [group for group in groups if group["values"]]


def history_detail_payload(wb, building: str) -> dict[str, Any]:
    plan = wb[PLAN_SHEET]
    plan_row = find_plan_row(plan, building)
    region = active_region(plan, plan_row)

    fiber_ws = wb[FIBER_SHEET]
    fiber_row = find_building_row(fiber_ws, building, 3)
    fiber = {
        "sheet": FIBER_SHEET,
        "row": fiber_row,
        "lastDate": display_value(fiber_ws.cell(fiber_row, 4).value),
        "values": [
            {
                "letter": get_column_letter(col),
                "label": norm(fiber_ws.cell(3, col).value),
                "value": display_value(fiber_ws.cell(fiber_row, col).value),
            }
            for col in range(8, 11)
        ],
    }

    measurement = None
    block = find_measure_block(wb, region, building)
    if block:
        sheet_name = block["sheet"]
        ws = wb[sheet_name]
        header_row = block["headerRow"]
        title = block.get("title") or building
        matches_building = True
        month_rows = []
        for row in range(header_row + 3, min(header_row + 15, ws.max_row + 1)):
            month_label = norm(ws.cell(row, 2).value)
            if not month_label:
                continue
            values = []
            for col in range(3, 31):
                value = ws.cell(row, col).value
                if value in (None, ""):
                    continue
                values.append(
                    {
                        "letter": get_column_letter(col),
                        "channel": norm(ws.cell(header_row, col).value),
                        "measure": norm(ws.cell(header_row + 1, col).value),
                        "kind": norm(ws.cell(header_row + 2, col).value),
                        "value": display_value(value),
                    }
                )
            groups = grouped_measure_values(ws, row, header_row)
            month_rows.append(
                {
                    "row": row,
                    "month": month_label,
                    "filledCount": len(values),
                    "values": values,
                    "groups": groups,
                    "location": display_value(ws.cell(row, 28).value),
                    "note": display_value(ws.cell(row, 29).value),
                    "inspector": display_value(ws.cell(row, 30).value),
                }
            )
        measurement = {
            "sheet": sheet_name,
            "sheetBuilding": title,
            "matchesBuilding": matches_building,
            "headerRow": header_row,
            "months": month_rows,
        }
    else:
        sheet_name = MEASURE_SHEETS.get(region)
        if sheet_name and sheet_name in wb.sheetnames:
            measurement = {
                "sheet": sheet_name,
                "sheetBuilding": None,
                "matchesBuilding": False,
                "headerRow": None,
                "months": [],
            }

    checklist = None
    sheet_name_for_building = find_sheet_name(wb, building)
    if sheet_name_for_building:
        ws = wb[sheet_name_for_building]
        items = []
        for row in range(5, min(ws.max_row, 24) + 1):
            item = norm(ws.cell(row, 5).value)
            result = display_value(ws.cell(row, 8).value)
            if item or result not in (None, ""):
                items.append(
                    {
                        "row": row,
                        "category": norm(ws.cell(row, 2).value),
                        "equipment": norm(ws.cell(row, 3).value),
                        "item": item,
                        "method": norm(ws.cell(row, 6).value),
                        "standard": norm(ws.cell(row, 7).value),
                        "result": result,
                    }
                )
        checklist = {"sheet": sheet_name_for_building, "items": items}

    return {
        "building": building,
        "region": region,
        "fiber": fiber,
        "measurement": measurement,
        "checklist": checklist,
    }


class Handler(BaseHTTPRequestHandler):
    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/" or path == "/index.html":
            self.send_file(ROOT / "index.html", "text/html; charset=utf-8")
        elif path == "/app.css":
            self.send_file(ROOT / "app.css", "text/css; charset=utf-8")
        elif path == "/app.js":
            self.send_file(ROOT / "app.js", "text/javascript; charset=utf-8")
        elif path == "/api/bootstrap":
            wb = load_workbook()
            now = today()
            self.send_json(
                {
                    "sectors": sectors_payload(wb),
                    "today": now.strftime("%Y-%m-%d"),
                    "month": now.month,
                    "dayLabel": f"{now.day:02d}일",
                    "workbook": str(WORKBOOK_PATH),
                }
            )
        elif path == "/api/history":
            wb = load_workbook()
            self.send_json({"ok": True, **plan_history_payload(wb)})
        elif path == "/api/history/detail":
            query = parse_qs(urlparse(self.path).query)
            building = norm((query.get("building") or [""])[0])
            if not building:
                raise ValueError("building 파라미터가 필요합니다.")
            wb = load_workbook()
            self.send_json({"ok": True, **history_detail_payload(wb, building)})
        elif path == "/download":
            ensure_workbook()
            self.send_file(
                WORKBOOK_PATH,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                download_name=WORKBOOK_PATH.name,
            )
        else:
            self.send_error(404)

    def do_POST(self):
        try:
            body = json.loads(self.rfile.read(int(self.headers.get("Content-Length", 0))) or "{}")
            if self.path == "/api/start":
                self.api_start(body)
            elif self.path == "/api/fiber/save":
                self.api_fiber_save(body)
            elif self.path == "/api/signal":
                self.api_signal(body)
            elif self.path == "/api/signal/save":
                self.api_signal_save(body)
            else:
                self.send_error(404)
        except Exception as exc:
            self.send_json({"ok": False, "error": str(exc)}, status=400)

    def api_start(self, body: dict[str, Any]) -> None:
        building = norm(body.get("building"))
        wb = load_workbook()
        plan = wb[PLAN_SHEET]
        now = today()
        plan_row = find_plan_row(plan, building)
        month_col = find_month_column(plan, now.month)
        region = active_region(plan, plan_row)
        plan.cell(plan_row, month_col).value = f"{now.day:02d}일"
        wb.save(WORKBOOK_PATH)
        self.send_json(
            {
                "ok": True,
                "building": building,
                "region": region,
                "planRow": plan_row,
                "planColumn": get_column_letter(month_col),
                "written": f"{now.day:02d}일",
                "measureSheet": MEASURE_SHEETS.get(region, ""),
            }
        )

    def api_fiber_save(self, body: dict[str, Any]) -> None:
        building = norm(body.get("building"))
        values = body.get("values", {})
        wb = load_workbook()
        ws = wb[FIBER_SHEET]
        row = find_building_row(ws, building, 3)
        changed = False
        for col in range(8, 11):
            key = get_column_letter(col)
            new_value = numeric_or_text(values.get(key, ""))
            if ws.cell(row, col).value != new_value:
                changed = True
            ws.cell(row, col).value = new_value
        if changed:
            ws.cell(row, 4).value = today().replace(tzinfo=None)
            ws.cell(row, 4).number_format = "yyyy-mm-dd"
        wb.save(WORKBOOK_PATH)
        self.send_json({"ok": True, "changed": changed, "row": row, "headers": fiber_headers(ws)})

    def api_signal(self, body: dict[str, Any]) -> None:
        building = norm(body.get("building"))
        step = int(body.get("step", 0))
        wb = load_workbook()
        plan = wb[PLAN_SHEET]
        plan_row = find_plan_row(plan, building)
        region = active_region(plan, plan_row)
        block = find_measure_block(wb, region, building)
        if not block:
            raise ValueError(f"{building}에 연결된 측정값 블록을 찾을 수 없습니다.")
        sheet_name = block["sheet"]
        header_row = block["headerRow"]
        ws = wb[sheet_name]
        month_row = find_month_row_in_block(ws, header_row, today().month)
        cols = SIGNAL_STEPS[step]
        values = {}
        previous = {}
        for col in cols:
            letter = get_column_letter(col)
            current = ws.cell(month_row, col).value
            prior = previous_nonblank_in_block(ws, header_row, month_row, col)
            previous[letter] = display_value(prior)
            values[letter] = display_value(current if current not in (None, "") else prior if col == 28 else current)
        self.send_json(
            {
                "ok": True,
                "sheet": sheet_name,
                "region": region,
                "monthRow": month_row,
                "headerRow": header_row,
                "step": step,
                "totalSteps": len(SIGNAL_STEPS),
                "headers": signal_headers(ws, cols, header_row),
                "values": values,
                "previous": previous,
            }
        )

    def api_signal_save(self, body: dict[str, Any]) -> None:
        building = norm(body.get("building"))
        step = int(body.get("step", 0))
        values = body.get("values", {})
        inspector = norm(body.get("inspector"))
        wb = load_workbook()
        plan = wb[PLAN_SHEET]
        plan_row = find_plan_row(plan, building)
        region = active_region(plan, plan_row)
        block = find_measure_block(wb, region, building)
        if not block:
            raise ValueError(f"{building}에 연결된 측정값 블록을 찾을 수 없습니다.")
        sheet_name = block["sheet"]
        header_row = block["headerRow"]
        ws = wb[sheet_name]
        month_row = find_month_row_in_block(ws, header_row, today().month)
        for col in SIGNAL_STEPS[step]:
            key = get_column_letter(col)
            raw = values.get(key, "")
            if col == 28 and raw == "":
                ws.cell(month_row, col).value = previous_nonblank_in_block(ws, header_row, month_row, col)
            elif col == 30 and raw == "" and inspector:
                ws.cell(month_row, col).value = inspector
            else:
                ws.cell(month_row, col).value = None if raw == "" else numeric_or_text(raw)
        wb.save(WORKBOOK_PATH)
        self.send_json({"ok": True, "sheet": sheet_name, "monthRow": month_row, "nextStep": step + 1})

    def send_json(self, data: dict[str, Any], status: int = 200) -> None:
        payload = json.dumps(data, ensure_ascii=False, default=str).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def send_file(self, path: Path, content_type: str, download_name: str | None = None) -> None:
        payload = path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(payload)))
        if download_name:
            self.send_header("Content-Disposition", f'attachment; filename="{download_name}"')
        self.end_headers()
        self.wfile.write(payload)

    def log_message(self, fmt: str, *args):
        return


if __name__ == "__main__":
    ensure_workbook()
    server = ThreadingHTTPServer(("127.0.0.1", 8765), Handler)
    print("CATV checker running at http://127.0.0.1:8765")
    server.serve_forever()
